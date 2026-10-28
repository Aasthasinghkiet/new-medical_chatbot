
"""Reusable document-loading and embedding helpers for the medical RAG chatbot.

Supports PDF, DOCX, CSV, XLSX, TXT and Markdown.
Preserves source metadata for retrieval and evaluation.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import List, Sequence

import pandas as pd
from docx import Document as WordDocument

from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".csv",
    ".xlsx",
    ".txt",
    ".md",
    ".markdown",
}


def _source_metadata(path: Path, **extra) -> dict:
    """Create metadata used to identify the original source."""
    return {
        "source": str(path),
        "source_name": path.name,
        **extra,
    }


def _clean(value) -> str:
    """Normalize empty and missing values."""
    if value is None:
        return ""

    value = str(value).strip()

    if value.lower() in {"nan", "none", "null"}:
        return ""

    return value


# --------------------------------------------------
# PDF loader
# --------------------------------------------------

def _load_pdf(path: Path) -> List[Document]:
    """Load a PDF while retaining page-number metadata."""
    documents = PyPDFLoader(str(path)).load()
    results = []

    for doc in documents:
        page = doc.metadata.get("page")

        doc.metadata.update(
            _source_metadata(path, page=page)
        )

        if doc.page_content and doc.page_content.strip():
            results.append(doc)

    return results


# --------------------------------------------------
# DOCX loader
# --------------------------------------------------

def _load_docx(path: Path) -> List[Document]:
    """Extract paragraphs and tables from a Word document."""
    word = WordDocument(str(path))
    parts = []

    for paragraph in word.paragraphs:
        value = paragraph.text.strip()

        if value:
            parts.append(value)

    for table_no, table in enumerate(word.tables, start=1):
        for row_no, row in enumerate(table.rows, start=1):
            cells = [
                cell.text.strip().replace("\n", " ")
                for cell in row.cells
            ]

            if any(cells):
                parts.append(
                    f"Table {table_no}, row {row_no}: "
                    + " | ".join(cells)
                )

    content = "\n".join(parts).strip()

    if not content:
        return []

    return [
        Document(
            page_content=content,
            metadata=_source_metadata(path),
        )
    ]


# --------------------------------------------------
# CSV and Excel row processing
# --------------------------------------------------

def _row_text(
    path: Path,
    sheet: str,
    row_number: int,
    row: pd.Series,
    columns: Sequence,
) -> str:
    """Convert a dataset row into readable text for retrieval."""

    values = [
        (str(col), _clean(row[col]))
        for col in columns
    ]

    values = [
        (col, value)
        for col, value in values
        if value
    ]

    if not values:
        return ""

    lookup = {
        col.lower().strip(): value
        for col, value in values
    }

    question = next(
        (
            value
            for key, value in lookup.items()
            if key in {"question", "questions", "query"}
        ),
        "",
    )

    answer = next(
        (
            value
            for key, value in lookup.items()
            if key in {"answer", "answers", "response"}
        ),
        "",
    )

    # Keep question-answer pairs together.
    if question and answer:
        excluded = {
            "question",
            "questions",
            "query",
            "answer",
            "answers",
            "response",
        }

        extras = [
            (key, value)
            for key, value in values
            if key.lower().strip() not in excluded
        ]

        text = (
            f"Medical Q&A dataset: {path.name}\n"
            f"Question: {question}\n"
            f"Answer: {answer}"
        )

        if extras:
            text += "\nAdditional metadata: " + "; ".join(
                f"{key}: {value}"
                for key, value in extras
            )

        return text

    return (
        f"Dataset: {path.name}; "
        f"sheet: {sheet}; row: {row_number}\n"
        + "; ".join(
            f"{key}: {value}"
            for key, value in values
        )
    )


def _load_tabular(path: Path) -> List[Document]:
    """Load CSV and Excel datasets into searchable documents."""
    documents: List[Document] = []

    rows_per_doc = max(
        1,
        int(os.getenv("ROWS_PER_DOCUMENT", "5")),
    )

    max_csv_rows = max(
        1,
        int(os.getenv("MAX_CSV_ROWS", "250000")),
    )

    if path.suffix.lower() == ".csv":
        frames = [
            (
                "CSV",
                pd.read_csv(
                    path,
                    dtype=str,
                    keep_default_na=False,
                    nrows=max_csv_rows,
                    low_memory=False,
                ),
            )
        ]
    else:
        frames = list(
            pd.read_excel(
                path,
                sheet_name=None,
                dtype=str,
            ).items()
        )

    for sheet_name, frame in frames:
        frame = frame.fillna("")
        columns = list(frame.columns)

        if frame.empty:
            documents.append(
                Document(
                    page_content=(
                        f"Dataset {path.name}, "
                        f"sheet {sheet_name}; columns: "
                        + ", ".join(map(str, columns))
                    ),
                    metadata=_source_metadata(
                        path,
                        sheet=str(sheet_name),
                    ),
                )
            )
            continue

        lower_columns = {
            str(col).lower().strip()
            for col in columns
        }

        is_qa = (
            bool(
                lower_columns.intersection(
                    {"question", "questions", "query"}
                )
            )
            and bool(
                lower_columns.intersection(
                    {"answer", "answers", "response"}
                )
            )
        )

        # Keep each Q&A record intact.
        step = 1 if is_qa else rows_per_doc

        for start in range(0, len(frame), step):
            batch = frame.iloc[start:start + step]
            records = []

            for offset, (_, row) in enumerate(batch.iterrows()):
                row_number = start + offset + 2

                value = _row_text(
                    path,
                    str(sheet_name),
                    row_number,
                    row,
                    columns,
                )

                if value:
                    records.append(value)

            if records:
                documents.append(
                    Document(
                        page_content="\n\n".join(records),
                        metadata=_source_metadata(
                            path,
                            sheet=str(sheet_name),
                            row=start + 2,
                            row_end=start + len(batch) + 1,
                        ),
                    )
                )

    return documents


# --------------------------------------------------
# Load one file
# --------------------------------------------------

def load_file(path: str | Path) -> List[Document]:
    """Load one supported file into LangChain Documents."""
    path = Path(path).expanduser().resolve()

    if not path.is_file():
        return []

    try:
        suffix = path.suffix.lower()

        if suffix == ".pdf":
            return _load_pdf(path)

        if suffix == ".docx":
            return _load_docx(path)

        if suffix in {".csv", ".xlsx"}:
            return _load_tabular(path)

        if suffix in {".txt", ".md", ".markdown"}:
            documents = TextLoader(
                str(path),
                encoding="utf-8",
                autodetect_encoding=True,
            ).load()

            for doc in documents:
                doc.metadata.update(_source_metadata(path))

            return [
                doc
                for doc in documents
                if doc.page_content.strip()
            ]

    except Exception as exc:
        print(
            f"WARNING: Could not read {path.name}: "
            f"{type(exc).__name__}: {exc}"
        )

    return []


# --------------------------------------------------
# Load all supported knowledge-base files
# --------------------------------------------------

def load_knowledge_files(
    data: str | Path,
) -> List[Document]:
    """Recursively load every supported file in a directory."""
    path = Path(data).expanduser()

    if not path.exists():
        raise FileNotFoundError(
            f"Knowledge path does not exist: {path}"
        )

    if path.is_file():
        documents = load_file(path)

        if not documents:
            print(
                f"WARNING: No text extracted from {path.name}. "
                "Scanned PDFs may require OCR."
            )

        return documents

    files = sorted(
        file_path
        for file_path in path.rglob("*")
        if (
            file_path.is_file()
            and file_path.suffix.lower() in SUPPORTED_EXTENSIONS
        )
    )

    documents: List[Document] = []

    if not files:
        print(f"WARNING: No supported files found in {path}")
        return documents

    for number, file_path in enumerate(files, start=1):
        loaded = load_file(file_path)
        documents.extend(loaded)

        print(
            f"[{number}/{len(files)}] "
            f"{file_path.name}: "
            f"loaded {len(loaded)} document units"
        )

    return documents


# --------------------------------------------------
# Backward-compatible PDF-only loader
# --------------------------------------------------

def load_pdf_file(
    data: str | Path,
) -> List[Document]:
    """Load PDF files only, retaining page metadata."""
    path = Path(data).expanduser()

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(
                f"load_pdf_file expects a PDF file, "
                f"got: {path.name}"
            )

        return _load_pdf(path.resolve())

    if not path.exists():
        raise FileNotFoundError(
            f"PDF directory does not exist: {path}"
        )

    files = sorted(
        file_path
        for file_path in path.rglob("*")
        if (
            file_path.is_file()
            and file_path.suffix.lower() == ".pdf"
        )
    )

    documents: List[Document] = []

    for pdf_path in files:
        loaded = _load_pdf(pdf_path)
        documents.extend(loaded)

        print(
            f"Loaded {len(loaded)} pages "
            f"from {pdf_path.name}"
        )

    return documents


# --------------------------------------------------
# Preserve minimal source metadata
# --------------------------------------------------

def filter_to_minimal_docs(
    docs: List[Document],
) -> List[Document]:
    """Keep source and location metadata for evaluation."""
    keep_keys = {
        "source",
        "source_name",
        "page",
        "sheet",
        "row",
        "row_end",
        "start_index",
    }

    minimal = []

    for doc in docs:
        metadata = {
            key: value
            for key, value in doc.metadata.items()
            if key in keep_keys and value is not None
        }

        minimal.append(
            Document(
                page_content=doc.page_content,
                metadata=metadata,
            )
        )

    return minimal


# --------------------------------------------------
# Text splitting
# --------------------------------------------------

def text_split(
    extracted_data: List[Document],
) -> List[Document]:
    """Split documents into overlapping chunks."""
    chunk_size = max(
        100,
        int(os.getenv("CHUNK_SIZE", "1000")),
    )

    chunk_overlap = max(
        0,
        min(
            int(os.getenv("CHUNK_OVERLAP", "150")),
            chunk_size - 1,
        ),
    )

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        add_start_index=True,
        separators=[
            "\n\n",
            "\n",
            ". ",
            "; ",
            ", ",
            " ",
            "",
        ],
    )

    return [
        doc
        for doc in splitter.split_documents(extracted_data)
        if doc.page_content.strip()
    ]


# --------------------------------------------------
# Hugging Face embeddings (local cache / offline)
# --------------------------------------------------

def download_hugging_face_embeddings(
) -> HuggingFaceEmbeddings:
    """Load the existing embedding model from the local cache.

    This function does not attempt to download files from Hugging Face.
    The selected model must already be available locally.
    """
    model_name = os.getenv(
        "EMBEDDING_MODEL",
        "sentence-transformers/all-MiniLM-L6-v2",
    )

    device = os.getenv("EMBEDDING_DEVICE", "cpu")

    batch_size = max(
        1,
        int(os.getenv("EMBEDDING_BATCH_SIZE", "32")),
    )

    return HuggingFaceEmbeddings(
        model_name=model_name,
        model_kwargs={
            "device": device,
            "local_files_only": True,
        },
        encode_kwargs={
            "normalize_embeddings": True,
            "batch_size": batch_size,
        },
    )
